"""Update a Vietnamese DOCX from a validated full ablation run.

Requires python-docx and matplotlib; never edits the input report in place.
"""
from __future__ import annotations

import argparse
import csv
from statistics import median
from collections import Counter
from zipfile import ZipFile
import json
import sys
from importlib.metadata import version
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from summarize_ablation import summarize, validate
from ablation_common import file_digest, read_jsonl


def rewrite(p, value):
    if p.runs:
        p.runs[0].text = value
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(value)


def table_values(t):
    return [[c.text for c in r.cells] for r in t.rows]


def generate_figures(rows, directory):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    directory.mkdir(exist_ok=True)
    replay = [r for r in rows if r['section'] == 'replay']
    labels = [r['variant'] for r in replay]
    fig, axes = plt.subplots(1, 2, figsize=(9, 3.6))
    for ax, key, title in zip(axes, ['total_node_ratio', 'total_dfs_time_ratio'],
                              ['Tổng node / baseline', 'Tổng thời gian DFS / baseline']):
        ax.bar(labels, [r[key] for r in replay], color=['#4477AA','#EE6677','#228833'],
               hatch=['', '//', '..'])
        ax.axhline(1, color='black', linewidth=.7, linestyle='--')
        ax.set_title(title)
        ax.tick_params(axis='x', rotation=18)
    fig.suptitle(f"Replay cùng {replay[0]['snapshots']} trạng thái khả thi; mỗi trạng thái có trọng số 1")
    fig.tight_layout()
    for extension in ['png','svg']:
        fig.savefig(directory / f'ablation_search.{extension}', dpi=300)
    plt.close(fig)
    info = [r for r in rows if r['section'] == 'info']
    presets = ['beginner', 'intermediate', 'expert']
    fig, ax = plt.subplots(figsize=(8, 4))
    for offset, variant, color, hatch in [(-.18,'baseline','#4477AA',''),(.18,'no_info','#EE6677','//')]:
        selected = [next(r for r in info if r['preset'] == p and r['variant'] == variant) for p in presets]
        rates = [100*r['win_rate'] for r in selected]
        ax.bar([i+offset for i in range(3)], rates, width=.36, label=variant, color=color, hatch=hatch,
               yerr=[[100*(r['win_rate']-r['ci_low']) for r in selected],
                     [100*(r['ci_high']-r['win_rate']) for r in selected]], capsize=4)
    ax.set_xticks(range(3), ['beginner (n=1000)','intermediate (n=1000)','expert (n=300)'])
    ax.set_ylabel('Tỉ lệ thắng (%)'); ax.set_ylim(0, 105)
    ax.set_title('INFO so với phá hòa tọa độ; Wilson 95% cho từng tỉ lệ')
    ax.legend(); fig.tight_layout()
    for extension in ['png','svg']:
        fig.savefig(directory / f'ablation_info.{extension}', dpi=300)
    plt.close(fig)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--template', type=Path, required=True)
    parser.add_argument('--data-dir', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args(argv)
    if args.out.exists():
        parser.error('output exists; choose a new path')
    manifest, configs, games, triples, excluded = validate(args.data_dir)
    if (manifest['games_override'] is not None or manifest['seed_start'] != 0 or excluded or
        sorted((c['preset'], c['games']) for c in configs.values()) !=
        [('beginner',1000),('expert',300),('intermediate',1000)]):
        parser.error('report requires the full planned sample and complete replay')
    expected_sizes={'beginner':(9,9,10),'intermediate':(16,16,40),'expert':(16,30,99)}
    for c in configs.values():
        if (c['rows'],c['cols'],c['mines']) != expected_sizes[c['preset']]:
            parser.error('report requires the planned preset dimensions/mines')
        if (c['first_click'],c['chord'],c['flood_fill'],c['opening_strategy'],c['include_outside']) != ('OPENING',True,'BFS','center',False):
            parser.error('report requires the default ablation configuration')
    rows = summarize(args.data_dir)
    generate_figures(rows, args.data_dir / 'figures')
    document = Document(args.template)
    original_tables = [table_values(t) for t in document.tables]
    original_images = [document.part.related_parts[s._inline.xpath('.//a:blip')[0].get(qn('r:embed'))].blob for s in document.inline_shapes]
    baseline = next(r for r in rows if r['section']=='replay' and r['variant']=='baseline')
    feasible = baseline['snapshots']
    summary_by_variant = {r['variant']:r for r in rows if r['section']=='replay'}
    nd, np = summary_by_variant['no_degree'], summary_by_variant['no_propagation']
    pairs = [r for r in rows if r['section']=='info_pair']
    replacements = {
        'Solver kết hợp các luật suy luận logic': 'Solver kết hợp suy luận logic, phân rã component, tìm kiếm và xác suất. So sánh solver hoàn chỉnh trên bàn 9x9 cho LV1 62.3%, LV2 79.1% và LV3 96.1%. Thí nghiệm ablation ở mục 4.2.7 đo riêng chi phí Degree Heuristic/propagation trên cùng frontier và ảnh hưởng INFO bằng cặp ván có cùng bàn ban đầu.',
        'Các luật LV1 và LV2 tạo kết luận tất định': 'Các luật LV1/LV2 tạo kết luận tất định khi thông tin và cờ nhất quán. Phân rã, propagation, pruning và Degree Heuristic phục vụ tìm kiếm; DFS đầy đủ phải giữ tập nghiệm. Tổng số mìn tạo trọng số xác suất; xác suất nhỏ nhất và INFO chọn nước đi. Mục 4.2.7 đối chiếu tập nghiệm/xác suất và chi phí của từng biến thể.',
        'Chưa có ablation cho Degree Heuristic': 'Đã triển khai DFSStats tùy chọn và chạy đầy đủ ablation tại mục 4.2.7. Kiểm thử có oracle brute-force độc lập cho component nhỏ và bàn nhỏ có tổng số mìn, đối chiếu nghiệm và Fraction. Những ca hữu hạn này cùng replay thực nghiệm không phải chứng minh thuật toán đúng với mọi đầu vào; cờ sai và trạng thái mâu thuẫn vẫn là hạn chế.',
        'Bổ sung DFSStats, thí nghiệm ablation': 'Mở rộng ablation và oracle brute-force sang nhiều cấu hình, component lớn hơn và môi trường khác; kiểm tra độ ổn định của hiệu ứng. DFSStats và replay cùng frontier đã có, nhưng không thay thế chứng minh tính đầy đủ hoặc đo hiệu năng phổ quát.',
        'Hệ thống hiện thực GUI, luật linh hoạt': 'Hệ thống hiện thực GUI, luật linh hoạt và phân tích headless. Bộ benchmark chính và ablation riêng đã được xác minh từ dữ liệu từng ván/frontier. Ablation tách chi phí DFS khỏi thay đổi chọn nước đi; kết quả chỉ áp dụng cho mẫu và cấu hình đã đo, chưa chứng minh hiệu năng phổ quát.',
        'Hai phương án thay thế được trình bày': 'Hai phương án thay thế được đo ở mục 4.2.5 và không đổi chiến lược mặc định. Constructor DfsSolver hỗ trợ include_outside=True; next_move truyền giá trị cấu hình này vào hàm xác suất, mặc định False.',
        'LV3 phá hòa bằng INFO': 'LV3 phá hòa bằng INFO (số ràng buộc gốc chứa ô), không dùng entropy hay mức giảm bất định kỳ vọng. Mục 4.2.7 so sánh INFO với phá hòa tọa độ trên cùng bàn ban đầu, chưa so với phá hòa ngẫu nhiên hoặc Information Gain. Solver chưa tối ưu trực tiếp xác suất thắng toàn ván hay số ô an toàn mở thêm.',
    }
    changed = []
    for p in document.paragraphs:
        for prefix, value in replacements.items():
            if p.text.startswith(prefix):
                changed.append(p.text); rewrite(p,value); break
    assert len(changed)==len(replacements), 'report template paragraphs changed'
    anchor = next(p for p in document.paragraphs if p.text.startswith('4.2 Kết quả'))
    def before(anchor, value, style='normal'):
        return anchor.insert_paragraph_before(value, style)
    before(anchor, f'Bộ ablation riêng {args.data_dir.name} chạy baseline 1.000/1.000/300 ván để thu frontier, cùng 1.000/1.000/300 cặp baseline/no_info (4.600 ván so sánh). Tổng cộng 6.900 ván; 2.300 ván thu frontier không được cộng vào cỡ mẫu INFO. Thu được {manifest["snapshots_unique"]} trạng thái duy nhất từ {manifest["snapshot_occurrences"]} lần xuất hiện. JSONL chỉ lưu PlayerView, ràng buộc gốc, cấu hình và seed/bước; không lưu bố trí mìn ẩn.')
    before(anchor, f'Replay baseline/no_degree/no_propagation dùng đúng cùng đầu vào, kiểm tra tập nghiệm từng component và xác suất Fraction bằng nhau. Mỗi biến thể warm-up một lần, đo {manifest["measured_repetitions"]} lần, lấy trung vị thời gian; thứ tự được xáo trộn bằng seed {manifest["shuffle_seed"]}, chạy tuần tự. Mỗi trạng thái duy nhất có trọng số 1. INFO ghép theo cấu hình/seed và fingerprint bàn sau nước đầu; khoảng Wilson mô tả từng tỉ lệ, McNemar exact hai phía đánh giá các cặp khác kết quả thắng.')
    before(anchor, 'DFSStats đếm mọi node (gốc, bị loại và nghiệm), node phân nhánh, node bị cắt một lần/node, phép gán forced không gồm phép gán nhánh, lượt quét ràng buộc, lượt propagation kể cả lượt cuối không đổi, độ sâu phân nhánh (gốc 0) và số nghiệm. Các bộ đếm cộng qua component/nước; độ sâu và kích thước component lấy max. DFS và tính xác suất đo riêng bằng perf_counter_ns; collect_stats=False vẫn là mặc định. no_propagation bỏ forced assignments nhưng giữ kiểm tra khả thi/cắt nhánh.')
    anchor = next(p for p in document.paragraphs if p.text.startswith('4.3 '))
    before(anchor, '4.2.7 Ablation Degree Heuristic, propagation và INFO','Heading 3')
    counts=Counter(s['input']['config']['config_id'] for s in read_jsonl(args.data_dir/'frontiers.jsonl'))
    before(anchor,'Số trạng thái duy nhất theo preset: '+', '.join(f"{c['preset']}={counts[cid]}" for cid,c in configs.items())+'. Bảng replay tổng hợp theo trạng thái, không cân bằng trọng số theo preset hoặc theo số ván.')
    before(anchor, f'Đối chiếu chính xác hoàn tất trên {baseline["matched_snapshots"]} frontier: {feasible} khả thi, {baseline["infeasible_snapshots"]} bất khả thi, {excluded} bị loại do không hoàn tất. Tất cả bộ ba khớp tập nghiệm và xác suất trong mọi lượt chạy. Kết quả không thay đổi hành vi mặc định; biến thể được cấu hình riêng, không monkeypatch, cache hay xấp xỉ.')
    def add_table(headers, values):
        t = document.add_table(rows=1, cols=len(headers))
        t.style = document.tables[0].style
        for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
        for values_row in values:
            for c,v in zip(t.add_row().cells,values_row): c.text=str(v)
        anchor._p.addprevious(t._tbl)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs: run.font.size=Pt(8)
        return t
    def f(x): return f'{x:.3f}'
    before(anchor,'Bảng 4.8 — Replay trên cùng frontier khả thi; mean/median/P95 theo node, thời gian ms.')
    add_table(['Biến thể','Node TB/trung vị/P95','Prune TB','Forced TB','DFS trung vị/P95 (ms)','Tỉ số tổng node/thời gian'],
              [[r['variant'], '/'.join(f(r[k]) for k in ['mean_nodes','median_nodes','p95_nodes']),f(r['mean_pruned_nodes']),f(r['mean_forced_assignments']),f(r['median_dfs_time_ms'])+'/'+f(r['p95_dfs_time_ms']),f(r['total_node_ratio'])+'/'+f(r['total_dfs_time_ratio'])] for r in rows if r['section']=='replay'])
    before(anchor, f'Bỏ Degree Heuristic có tỉ số tổng node {nd["total_node_ratio"]:.3f} và tổng thời gian DFS {nd["total_dfs_time_ratio"]:.3f} so với baseline. Bỏ propagation nhưng vẫn pruning có các tỉ số tương ứng {np["total_node_ratio"]:.3f} và {np["total_dfs_time_ratio"]:.3f}. Đây là tỉ số tổng trên cùng tập frontier, không phải trung bình các tỉ số hoặc đo trên các quỹ đạo chơi khác nhau. Thời gian rất ngắn có nhiễu đo và phụ thuộc máy; instrumentation bật cho cả ba biến thể.')
    for name, caption in [('ablation_search','Hình 4.3 — Tỉ số tổng node và thời gian DFS trên cùng frontier.'),('ablation_info','Hình 4.4 — Tỉ lệ thắng INFO và phá hòa tọa độ; Wilson 95%.')]:
        before(anchor,caption)
        before(anchor,'').add_run().add_picture(str(args.data_dir/'figures'/f'{name}.png'),width=Inches(5.8))
    before(anchor,'Bảng 4.9 — Ván baseline/no_info theo cặp. Thắng/thua/stalled và CI Wilson 95%.')
    add_table(['Preset/biến thể','n','Thắng/thua/stalled','Tỉ lệ thắng (CI 95%)','Nước TB','Đoán TB'],
              [[r['preset']+'/'+r['variant'],r['games'],f"{r['wins']}/{r['losses']}/{r['stalled']}",f"{100*r['win_rate']:.1f}% ({100*r['ci_low']:.1f}–{100*r['ci_high']:.1f})",f(r['mean_moves']),f(r['mean_guesses'])] for r in rows if r['section']=='info'])
    before(anchor,'Bảng 4.10 — Kiểm định thắng theo cặp; b=chỉ baseline thắng, c=chỉ no_info thắng; Δ=no_info−baseline (điểm phần trăm).')
    add_table(['Preset','Số cặp','b','c','Δ (đpt)','McNemar exact p'],[[r['preset'],r['games'],r['b'],r['c'],f(r['win_difference_pp']),f"{r['mcnemar_exact_p']:.6g}"] for r in pairs])
    nominal=[r['preset'] for r in pairs if r['mcnemar_exact_p']<.05]
    adjusted=[r['preset'] for r in pairs if r['mcnemar_exact_p']<.05/len(pairs)]
    before(anchor,f'Ở ngưỡng 0,05 chưa hiệu chỉnh, các preset có p nhỏ hơn ngưỡng là: {", ".join(nominal) or "không có"}. Nếu kiểm soát ba so sánh bằng Bonferroni (0,05/3 ≈ 0,0167), các preset đạt ngưỡng là: {", ".join(adjusted) or "không có"}. Vì vậy cần đọc kết quả theo từng preset và phương án đa kiểm định, không kết luận INFO luôn cải thiện tỉ lệ thắng.')
    before(anchor,'McNemar dùng xác suất nhị thức p₀=0,5 trên b+c cặp bất đồng, hai phía, chưa hiệu chỉnh đa kiểm định. Không suy ra tương đương từ p không có ý nghĩa hoặc khoảng Wilson chồng lấp. Chỉ so INFO với tọa độ trên cấu hình đã đo; node toàn ván của no_info không tách được tác động Degree/propagation vì quỹ đạo có thể thay đổi. Các frontier được thu từ baseline nên mẫu không đại diện mọi trạng thái có thể gặp.')
    stalls=sum(r['stalled'] for r in rows if r['section']=='info')
    before(anchor,f'Số ván stalled trong mẫu INFO: {stalls}; dữ liệu không chứa lỗi thực thi bị ghi nhận. Kết quả đo riêng trong ablation không được ghép với thời gian toàn ván của benchmark chính hoặc dùng để suy diễn lợi ích entropy. Oracle brute-force độc lập chỉ bao phủ các trường hợp nhỏ trong kiểm thử; replay giữa các biến thể dùng chung hàm tính xác suất nên không phải oracle độc lập cho toàn bộ mẫu thực tế.')
    raw_replay=list(csv.DictReader((args.data_dir/'ablation_frontiers.csv').open()))
    empty=sum(r['variant']=='baseline' and int(r['components_count'])==0 for r in raw_replay)
    before(anchor,f'Trong mẫu có {empty} trạng thái không có component frontier (đường fallback của LV3); bộ đếm DFS bằng 0. Các trạng thái này vẫn được ghi và đối chiếu đầu vào/xác suất, nằm trong thống kê trạng thái khả thi khi có cấu hình mìn toàn cục hợp lệ. Không coi chúng là một component DFS có biến.')
    max_depth=max(int(r['max_depth']) for r in raw_replay)
    largest=max(int(r['largest_component_variables']) for r in raw_replay)
    probability_times=[float(r['probability_time_ms']) for r in raw_replay if r['variant']=='baseline' and r['feasibility']=='feasible']
    before(anchor,f'Component lớn nhất có {largest} biến; độ sâu phân nhánh lớn nhất qua các biến thể là {max_depth} (gốc 0). Trung vị thời gian tính xác suất baseline riêng là {median(probability_times):.4f} ms/frontier khả thi, không cộng vào thời gian DFS. Lượt quét ràng buộc và propagation được lưu trong ablation_frontiers.csv; no_propagation có forced assignments và propagation_passes bằng 0.')
    for r in pairs:
        before(anchor,f"Trên {r['preset']}, thay INFO bằng tọa độ làm tỉ lệ thắng thay đổi {r['win_difference_pp']:+.2f} điểm phần trăm, với b={r['b']}, c={r['c']} và p exact={r['mcnemar_exact_p']:.6g}. Đây là chênh lệch quan sát trong các cặp đã đo; không diễn giải p lớn như chứng minh hai chiến lược tương đương.")
    before(anchor,'Các phân vị dùng nội suy tại (n−1)×q. Bảng chi tiết theo preset, counter/độ sâu/component và thời gian xác suất nằm trong CSV thô; bảng trên tổng hợp frontier khả thi của cả ba preset, không gán một trạng thái nhiều lần theo occurrences.')
    document.add_paragraph('Tái lập ablation đầy đủ (giữ mặc định, seed bắt đầu 0; cần matplotlib và python-docx để tạo hình/báo cáo):')
    for command in ['python docs/report/scripts/run_ablation.py --out-dir docs/report/data/ablation_new',
                    'python docs/report/scripts/summarize_ablation.py --data-dir docs/report/data/ablation_new']:
        document.add_paragraph(command)
    document.add_paragraph(f'Dữ liệu ablation: {args.data_dir.as_posix()}. Nguồn lúc đo: {manifest["git_commit"]}; manifest, source_snapshot.zip và checksum xác định chính xác mã đã chạy. Mỗi cấu hình mặc định có beginner/intermediate 1.000 và expert 300 cặp; không dùng smoke để kết luận. Chỉ tạo bảng khi toàn bộ ma trận, fingerprint và replay đã xác minh; lần bị cắt hoặc sai nghiệm không được coi là hoàn tất.')
    settings=document.settings.element
    field=settings.find(qn('w:updateFields'))
    if field is None:
        field=OxmlElement('w:updateFields');settings.append(field)
    field.set(qn('w:val'),'true')
    args.out.parent.mkdir(parents=True,exist_ok=True)
    document.save(args.out)
    checked=Document(args.out)
    assert [table_values(t) for t in checked.tables if table_values(t) in original_tables]==original_tables
    checked_blobs=[checked.part.related_parts[s._inline.xpath('.//a:blip')[0].get(qn('r:embed'))].blob for s in checked.inline_shapes]
    assert all(b in checked_blobs for b in original_images)
    assert all((args.data_dir/'figures'/f'{name}.png').read_bytes() in checked_blobs for name in ['ablation_search','ablation_info'])
    assert len(checked.tables)==len(original_tables)+3
    assert len(checked.inline_shapes)==len(original_images)+2
    assert any('TOC ' in v for v in checked._element.xpath('//w:instrText/text()'))
    with ZipFile(args.out) as archive:
        assert archive.testzip() is None
    assert all(old not in [p.text for p in checked.paragraphs] for old in changed)
    audit=dict(command=sys.argv,environment={'python':sys.version,'python-docx':version('python-docx'),'matplotlib':version('matplotlib')},template_sha256=file_digest(args.template),report_sha256=file_digest(args.out),
               manifest_sha256=file_digest(args.data_dir/'manifest.json'),summary_sha256=file_digest(args.data_dir/'ablation_summary.csv'),
               script_sha256=file_digest(Path(__file__)),
               figure_sha256={p.name:file_digest(p) for p in sorted((args.data_dir/'figures').iterdir())},
               main_tables_preserved=True,original_images_preserved=True,
               docx_zip_valid=True,new_figures_match_png=True,word_toc_field_preserved=True,
               exact_replay_verified=True,paired_board_fingerprints_verified=True,excluded_snapshots=excluded,
               snapshots=baseline['matched_snapshots'],new_tables=3,new_figures=2,pagination_rendered=False,
               updated_paragraphs=len(changed))
    args.out.with_suffix('.audit.json').write_text(json.dumps(audit,indent=2)+'\n')
    print(json.dumps(audit,indent=2))
    return 0


if __name__=='__main__':
    raise SystemExit(main())
